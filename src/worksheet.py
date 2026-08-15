from openpyxl import load_workbook
from openpyxl.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from src.paths import FARMS_SHEET_PATH

workbook = load_workbook(FARMS_SHEET_PATH)
_sheet = workbook.active


def get_sheet() -> Worksheet:
    assert isinstance(_sheet, Worksheet), "No active sheet in workbook"
    return _sheet

def save_workbook() -> None:
    workbook.save(FARMS_SHEET_PATH)


def get_column(column: str) -> list:
    for col in get_sheet().iter_cols(values_only=True):
        if col[0] == column:
            return list(col[1:])
    raise ValueError(f"Column {column} not found in worksheet")


def get_row(row: str) -> list[Cell]:
    for row_values in get_sheet().iter_rows():
        if row_values[0].value == row:
            return list(row_values) # type: ignore
    raise ValueError(f"Row {row} not found in worksheet. possible row names: {get_sheet()['A']}")


keys = [key.value for key in get_sheet()[1] if isinstance(key.value, str)]
